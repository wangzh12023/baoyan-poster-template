"""版式：把 poster.yaml 的内容按固定版式排到底稿上，并按一款配色上色。

底稿 template/base.pptx 只有背景渐变和 5 枚装饰渐变圆，其余文字、线条、logo、二维码都在这里现画。
坐标单位 pt（1pt = 12700 EMU），海报 768 × 1152 pt（2:3，打印可等比放大到 A3 / A2 等）。"""
import io
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI

from . import fonts as F
from . import logo as LOGO
from . import qr as QR
from .palettes import tint

TEMPLATE = Path(__file__).resolve().parent / "template" / "base.pptx"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_SVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
NSMAP = f'xmlns:p="{NS_P}" xmlns:a="{NS_A}"'
PT = 12700

# ---------------- 版式（pt） ----------------
L, R, GUTTER = 43.5, 726, 40
COL_W = (R - L - GUTTER) / 2
COL_X = [L, L + COL_W + GUTTER]
YEAR_X, YEAR_Y, YEAR_W, YEAR_H, YEAR_RULE_Y, YEAR_RULE_END = 41.25, 42.75, 225, 90, 87, 468.75
TITLE_X, TITLE_Y, TITLE_W, TITLE_H = 37.5, 129.75, 477.75, 107.25
SUB_Y = [(242.25, 24, 470.25), (267, 23.25, 453.75)]                 # 英文副标题两行：y, h, w
DATE_Y, DATE_H = 23.25, 66.75
TIME_Y, TIME_H, TIME_W = 92.2, 33.75, 220.5
VENUE_Y, VENUE_H, VENUE_W = 130.35, 35.25, 209.25
HEADER_TOP, HEADER_H, HEADER_RULE = 336, 34, 377
CARD_TOP0 = 394
NAME_H, INST_OFF, INST_H, TOPIC_OFF, TOPIC_H = 27, 29, 23, 58, 25
CARD_H = TOPIC_OFF + TOPIC_H
FOOT_END, FOOT_H, FOOT_GAP = 903, 20, 26
CARD_SPAN = FOOT_END - FOOT_H - FOOT_GAP - CARD_H - CARD_TOP0      # 第一张卡片顶 → 最后一张卡片顶
PITCH_MAX = CARD_SPAN / 3                                           # 每场 4 位嘉宾时正好排满
DOT, DOT_GAP = 6.75, 9
BAR_RULE = 928
BAR_TOP = BAR_RULE + 22          # 底栏各栏小标签顶端
VAL_OFF = 21                     # 标签 → 内容
QR_SIZE = 116
QR_X = R - QR_SIZE
ORG_X, DIV1_X, DIV2_X = 352, 330, QR_X - 22
LOGO_H, LOGO_MAX_W = 68, DIV2_X - 22 - ORG_X
HAIRLINE = 1                      # QuickLook 会丢掉 <1pt 的形状

NAME_SZ, INST_SZ, TOPIC_SZ, SLOT_SZ, KICK_SZ = 21, 15.5, 19, 13.5, 14
TITLE_SZ = 90


def emu(pt):
    return int(round(pt * PT))


def esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def run(text, font, sz, color, b=False, i=False, spc=0):
    spc_attr = f' spc="{round(spc * 100)}"' if spc else ""
    return (f'<a:r><a:rPr lang="zh-CN" altLang="en-US" sz="{round(sz * 100)}" b="{int(b)}" i="{int(i)}"{spc_attr} '
            f'dirty="0"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'<a:latin typeface="{esc(font[0])}"/><a:ea typeface="{esc(font[1])}"/><a:cs typeface="{esc(font[0])}"/>'
            f'</a:rPr><a:t>{esc(text)}</a:t></a:r>')


def para(runs, algn="l"):
    return f'<a:p><a:pPr algn="{algn}"></a:pPr>{"".join(runs)}</a:p>'


def txbody(paras, anchor="ctr"):
    return etree.fromstring(
        f'<p:txBody {NSMAP}><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="{anchor}">'
        f'<a:noAutofit/></a:bodyPr><a:lstStyle/>{"".join(paras)}</p:txBody>')


def sp_xml(sid, name, x, y, w, h, fill=None, paras=None, geom="rect"):
    fill_xml = f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill>' if fill else "<a:noFill/>"
    el = etree.fromstring(
        f'<p:sp {NSMAP}><p:nvSpPr><p:cNvPr id="{sid}" name="{esc(name)}"/><p:cNvSpPr><a:spLocks noGrp="1"/>'
        f'</p:cNvSpPr><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="{emu(x)}" y="{emu(y)}"/>'
        f'<a:ext cx="{emu(w)}" cy="{emu(h)}"/></a:xfrm><a:prstGeom prst="{geom}"><a:avLst/></a:prstGeom>'
        f'{fill_xml}<a:ln w="0"><a:noFill/></a:ln></p:spPr></p:sp>')
    if paras is not None:
        el.append(txbody(paras))
    return el


_said = set()


def warn(msg):
    """同一条提示只说一次（54 款配色共用同一份内容）"""
    if msg not in _said:
        _said.add(msg)
        print(f"  ⚠ {msg}", file=sys.stderr)


def fit(text_w, max_w, sz, min_sz, what):
    """字太长时按比例缩小字号（不低于 min_sz），返回新字号"""
    if text_w <= max_w:
        return sz
    new = max(min_sz, sz * max_w / text_w)
    if text_w * new / sz > max_w + 0.5:
        warn(f"{what}太长，缩到 {new:.1f}pt 仍放不下，建议删减文字")
    else:
        warn(f"{what}较长，已自动缩小到 {new:.1f}pt")
    return new


class Slide:
    """往幻灯片上加形状的小工具"""

    def __init__(self, prs):
        self.prs, self.slide = prs, prs.slides[0]
        self.tree = self.slide.shapes._spTree
        self.next_id = max(int(e.get("id")) for e in self.tree.iter(f"{{{NS_P}}}cNvPr"))

    def add(self, name, x, y, w, h, fill=None, paras=None, geom="rect"):
        self.next_id += 1
        el = sp_xml(self.next_id, name, x, y, w, h, fill, paras, geom)
        ext = self.tree.find(f"{{{NS_P}}}extLst")
        ext.addprevious(el) if ext is not None else self.tree.append(el)
        return el

    def picture(self, name, png, svg, x, y, w, h, svg_name):
        pic = self.slide.shapes.add_picture(io.BytesIO(png), emu(x), emu(y), emu(w), emu(h))
        pic.name = name
        if svg:                                            # 矢量副本：PowerPoint 2016+ / Keynote 优先用它
            part = Part(PackURI(f"/ppt/media/{svg_name}.svg"), "image/svg+xml", self.slide.part.package, svg)
            rid = self.slide.part.relate_to(part, RT.IMAGE)
            pic._element.find(f".//{{{NS_A}}}blip").append(etree.fromstring(
                f'<a:extLst xmlns:a="{NS_A}"><a:ext uri="{{96DAC541-7B7A-43D3-8B79-37D633B846F1}}">'
                f'<asvg:svgBlip xmlns:asvg="{NS_SVG}" xmlns:r="{NS_R}" r:embed="{rid}"/></a:ext></a:extLst>'))
        return pic


def build(cfg, p):
    """cfg：config.load() 的结果；p：palettes 里的一款配色。返回 Presentation"""
    INK, BODY, MUTED, ACC, HAIR, POP, TOPIC = (p[k] for k in ("ink", "body", "muted", "accent", "hair", "pop", "topic"))
    FT = F.FONTS
    SERIF_BLACK, SERIF, SANS, SANS_MED, SANS_SEMI = (FT[k] for k in ("serif_black", "serif", "sans", "sans_medium",
                                                                     "sans_semibold"))
    DISPLAY, LAT_MED, LAT_DEMI = FT["display"], FT["latin_medium"], FT["latin_demibold"]
    W = F.width

    prs = Presentation(TEMPLATE)
    s = Slide(prs)
    add = s.add

    # ---- 背景渐变 + 装饰渐变圆：色标按配色换色相 ----
    grad = s.slide._element.find(f".//{{{NS_P}}}bgPr")[0]
    for c, v in zip(grad.iter(f"{{{NS_A}}}srgbClr"), p["bg"]):
        c.set("val", v)
    for shape in s.slide.shapes:
        if shape.name.startswith("装饰圆"):
            for c in shape._element.iter(f"{{{NS_A}}}srgbClr"):
                c.set("val", tint(c.get("val"), p["hue"], p["drop"], p["chroma"], p["soft"], p["mul"]))

    # ---- 右上：日期 + 星期一行，时间 / 地点两行；整组右边缘对齐版心 ----
    date_w = W(cfg["date"], DISPLAY, 56)
    if cfg["weekday"]:
        date_w += W("  ", DISPLAY, 20) + W(cfg["weekday"], DISPLAY, 21)
    block_w = max(date_w, W(cfg["time"], LAT_MED, 21) + 2, W(cfg["venue"], LAT_MED, 21) + 2)
    bx = R - block_w
    date_runs = [run(cfg["date"], DISPLAY, 56, INK)]
    if cfg["weekday"]:
        date_runs += [run("  ", DISPLAY, 20, INK), run(cfg["weekday"], DISPLAY, 21, ACC, i=True)]
    add("日期", bx, DATE_Y, date_w + 6, DATE_H, paras=[para(date_runs)])
    for name, text, y, h, w in (("时间", cfg["time"], TIME_Y, TIME_H, TIME_W),
                                ("地点", cfg["venue"], VENUE_Y, VENUE_H, VENUE_W)):
        if text:
            add(name, bx + 2, y, max(w, block_w + 4), h, paras=[para([run(text, LAT_MED, 21, INK)])])

    # ---- 左上：年份 + 横线，主标题，英文副标题（第一行拉开字距与主标题等宽） ----
    if cfg["year"]:
        add("年份", YEAR_X, YEAR_Y, YEAR_W, YEAR_H, paras=[para([run(cfg["year"], DISPLAY, 84, INK)])])
        x = YEAR_X + W(cfg["year"], DISPLAY, 84) + 22
        if YEAR_RULE_END - x > 30:
            add("年份右横线", x, YEAR_RULE_Y, YEAR_RULE_END - x, 1.5, fill=POP)
    title_max = bx - 16 - TITLE_X
    title_sz = fit(W(cfg["title"], SERIF_BLACK, TITLE_SZ), title_max, TITLE_SZ, 40, "主标题")
    title_w = W(cfg["title"], SERIF_BLACK, title_sz)
    sink = F.base_off((SERIF_BLACK, TITLE_SZ)) - F.base_off((SERIF_BLACK, title_sz))   # 缩小后基线不动
    add("主标题", TITLE_X, TITLE_Y + sink, max(TITLE_W, title_w + 10), TITLE_H,
        paras=[para([run(cfg["title"], SERIF_BLACK, title_sz, INK)])])
    if cfg["subtitle"]:
        first = cfg["subtitle"][0]
        track = max(0.0, (title_w - 4 - W(first, LAT_DEMI, 12)) / len(first))
        for text, (y, h, w) in zip(cfg["subtitle"], SUB_Y):
            add("英文副标题", TITLE_X + 3, y, w, h, paras=[para([run(text, LAT_DEMI, 12, BODY, spc=track)])])

    # ---- 嘉宾区：两栏，每栏「编号 + 场次名 + 时间」+ 粗标题线，嘉宾之间细分隔线 ----
    n_max = max(len(sess["speakers"]) for sess in cfg["sessions"])
    pitch = min(PITCH_MAX, CARD_SPAN / (n_max - 1)) if n_max > 1 else PITCH_MAX
    head_off = F.base_off((DISPLAY, 30), (SERIF, 17))
    name_base = F.base_off((SERIF_BLACK, NAME_SZ))
    for ci, sess in enumerate(cfg["sessions"]):
        x = COL_X[ci]
        label = f"第{ci + 1}场"
        add(f"{label}·标签", x, HEADER_TOP, COL_W, HEADER_H,
            paras=[para([run(f"0{ci + 1}", DISPLAY, 30, ACC), run("   ", SERIF, 17, INK),
                         run(sess["name"], SERIF, 17, INK, b=True)])])
        if sess["time"]:
            add(f"{label}·时间", x, HEADER_TOP + head_off - F.base_off((LAT_MED, 15)), COL_W, HEADER_H,
                paras=[para([run(sess["time"], LAT_MED, 15, BODY)], algn="r")])
        add(f"{label}·标题线", x, HEADER_RULE, COL_W, 1, fill=INK)

        for i, sp in enumerate(sess["speakers"]):
            top = CARD_TOP0 + i * pitch
            who = sp["name"]
            if i:
                add(f"{who}·上分隔线", x, top - (pitch - CARD_H) / 2, COL_W, HAIRLINE, fill=HAIR)
            add(f"{who}·姓名", x, top, COL_W, NAME_H, paras=[para([run(who, SERIF_BLACK, NAME_SZ, INK)])])
            used = W(who, SERIF_BLACK, NAME_SZ)
            if sp["tag"]:
                kw = W(sp["tag"], SANS_SEMI, KICK_SZ, 0.6) + 4
                add(f"{who}·标签", x + used + 9, top + name_base - F.base_off((SANS_SEMI, KICK_SZ)), kw, NAME_H,
                    paras=[para([run(sp["tag"], SANS_SEMI, KICK_SZ, ACC, spc=0.6)])])
                used += 9 + kw
            if sp["time"]:
                add(f"{who}·时间", x, top + name_base - F.base_off((LAT_MED, SLOT_SZ)), COL_W, NAME_H,
                    paras=[para([run(sp["time"], LAT_MED, SLOT_SZ, MUTED)], algn="r")])
                if used + 12 + W(sp["time"], LAT_MED, SLOT_SZ) > COL_W:
                    warn(f"「{who}」的姓名 + 标签和右侧时间挤在一起了，建议缩短标签")
            if sp["school"] or sp["dept"]:
                dept = ("  " + sp["dept"]) if sp["school"] and sp["dept"] else sp["dept"]
                sz = fit(W(sp["school"], SANS_MED, INST_SZ) + W(dept, SANS, INST_SZ), COL_W - 4, INST_SZ, 12,
                         f"「{who}」的院校")
                add(f"{who}·去向", x, top + INST_OFF, COL_W, INST_H,
                    paras=[para([run(sp["school"], SANS_MED, sz, BODY)] * bool(sp["school"]) +
                                [run(dept, SANS, sz, BODY)] * bool(dept))])
            if sp["topic"]:
                sz = fit(W(sp["topic"], SANS_SEMI, TOPIC_SZ), COL_W - 4, TOPIC_SZ, 14, f"「{who}」的分享主题")
                add(f"{who}·分享主题", x, top + TOPIC_OFF, COL_W, TOPIC_H,
                    paras=[para([run(sp["topic"], SANS_SEMI, sz, TOPIC)])])

        fn = sess["footnote"]
        if fn:
            foot_top = FOOT_END - FOOT_H
            add(f"{label}·末尾分隔线", x, foot_top - FOOT_GAP / 2, COL_W, HAIRLINE, fill=HAIR)
            add(f"{label}·末尾圆点", x, foot_top + (FOOT_H - DOT) / 2, DOT, DOT, fill=POP, geom="ellipse")
            runs = ([run(fn["time"], LAT_MED, 12.5, MUTED)] if fn["time"] else []) + \
                   [run(("   " if fn["time"] else "") + fn["text"], SANS, 12.5, MUTED)]
            add(f"{label}·末尾说明", x + DOT + DOT_GAP, foot_top, COL_W - DOT - DOT_GAP, FOOT_H, paras=[para(runs)])

    # ---- 底栏：时间·地点 | 主办 logo | 二维码，三栏统一「小标签 + 内容」 ----
    add("底栏上横线", L, BAR_RULE, R - L, 1, fill=INK)

    def label(name, pair, x, y, w=220):
        cn, en = pair
        add(name, x, y, w, 16, paras=[para([run(cn, SANS_MED, 10.5, MUTED)] +
                                           ([run("   " + en, LAT_DEMI, 8.5, MUTED, spc=1.2)] if en else []))])

    y0, bar_end = BAR_TOP, BAR_TOP + VAL_OFF + QR_SIZE + 22
    ft, lab = cfg["footer"], cfg["labels"]
    left_w = DIV1_X - L - 12
    if ft["date"] or ft["time"]:
        label("底栏·时间标签", lab["time"], L, y0)
        sep = "   " if ft["date"] and ft["time"] else ""
        sz = fit(W(ft["date"], SANS_MED, 16) + W(sep + ft["time"], LAT_DEMI, 16), left_w, 16, 12, "底栏日期时间")
        add("底栏·日期时间", L, y0 + VAL_OFF - 2, 270, 24,
            paras=[para([run(ft["date"], SANS_MED, sz, INK), run(sep + ft["time"], LAT_DEMI, sz, INK)])])
    if ft["venue"]:
        label("底栏·地点标签", lab["venue"], L, y0 + 62)
        sz = fit(W(ft["venue"], LAT_DEMI, 16), left_w, 16, 12, "底栏地点")
        add("底栏·地点", L, y0 + 62 + VAL_OFF - 2, 270, 24, paras=[para([run(ft["venue"], LAT_DEMI, sz, INK)])])
    if ft["note"]:
        add("底栏·发起人", L, bar_end - 16, 280, 16, paras=[para([run(ft["note"], SANS, 11, MUTED)])])

    lg = cfg["logo"]
    if lg:
        add("底栏·分隔线一", DIV1_X, y0, HAIRLINE, bar_end - y0, fill=HAIR)
        label("底栏·主办标签", lab["organizer"], ORG_X, y0)
        if lg["file"]:
            png, svg, aspect = LOGO.load(lg["file"], LOGO.target_color(lg["color"], p))
            h = min(LOGO_H, LOGO_MAX_W / aspect)
            s.picture("主办单位logo", png, svg, ORG_X, y0 + VAL_OFF + 3, h * aspect, h, "logo")
        else:
            for li, line in enumerate(lg["text"][:3]):
                add("主办单位", ORG_X, y0 + VAL_OFF - 2 + li * 26, LOGO_MAX_W, 24,
                    paras=[para([run(line, SANS_MED, 15, INK)])])

    q = cfg["qr"]
    if q:
        add("底栏·分隔线二", DIV2_X, y0, HAIRLINE, bar_end - y0, fill=HAIR)
        label("二维码标签", q["title"], QR_X, y0, QR_SIZE)
        png, svg = QR.make(q, INK)
        s.picture("二维码", png, svg, QR_X, y0 + VAL_OFF, QR_SIZE, QR_SIZE, "qr")
        if q["id"] or q["id_label"]:
            add("二维码下方文字", QR_X, bar_end - 16, QR_SIZE, 16,
                paras=[para(([run(q["id_label"] + "  ", SANS, 10.5, MUTED)] if q["id_label"] else []) +
                            ([run(q["id"], LAT_MED, 10.5, INK, spc=0.4)] if q["id"] else []), algn="ctr")])

    cp = prs.core_properties
    cp.title = f"{cfg['year']} {cfg['title']}".strip()
    cp.subject = f"{p['key']} {p['name']}"
    cp.author = cp.last_modified_by = ""
    cp.comments = "Made with baoyan-poster-template"
    return prs
